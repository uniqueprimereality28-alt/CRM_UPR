"""
Attendance Export Engine for Unique Prime Reality CRM.
Generates:
1. Multi-sheet Excel workbook (.xlsx) with:
   - Sheet 1: Monthly Summary (S.No, Name, Role, Days Present, Worked, Overtime, Late Markings)
   - Sheet 2: Day-wise Matrix (1st to 31st with P / A / OFF)
   - Sheets 3+: Individual Employee Tabs with date-by-date In/Out timings in IST.
2. Formatted PDF statement (.pdf) with company branding, KPI cards, and summary/punch tables.
3. Role masking: displays Vrinda Aggarwal's role as 'Technical Head' to preserve corporate harmony.
"""

import io
import re
from datetime import datetime, timedelta, timezone
from typing import List, Dict, Any, Optional

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT


def mask_role(name: str, role: str) -> str:
    """Masks Superadmin role as 'Technical Head' for corporate presentation."""
    n = (name or "").strip().lower()
    r = (role or "").strip().lower()
    if "vranda" in n or "vrinda" in n or r == "superadmin":
        return "Technical Head"
    if r == "admin":
        return "Admin"
    if r in ("team_lead", "team lead", "tl"):
        return "Team Lead"
    if r == "sales":
        return "Sales"
    return role.replace("_", " ").title() if role else "-"


def sec_to_hm(seconds: int) -> str:
    """Converts seconds into 'Xh Ym' (e.g. 115h 00m)."""
    if not seconds or seconds <= 0:
        return "0h 00m"
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    return f"{h}h {m:02d}m"


def to_ist_str(iso_str: Optional[str]) -> str:
    """Converts a UTC ISO timestamp to IST (UTC+5:30) formatted string '11:12 AM'."""
    if not iso_str:
        return "—"
    try:
        dt = datetime.fromisoformat(iso_str)
        ist_dt = dt.astimezone(timezone(timedelta(hours=5, minutes=30)))
        return ist_dt.strftime("%I:%M %p")
    except Exception:
        return iso_str


def sanitize_sheet_title(title: str) -> str:
    """Sanitizes sheet titles to valid Excel names (max 30 chars, no illegal characters)."""
    cleaned = re.sub(r'[\\/*?:\[\]]', '', title or "Sheet")[:30].strip()
    return cleaned or "Sheet"


def build_attendance_report_data(
    attendance_records: List[Dict[str, Any]],
    users: List[Dict[str, Any]],
    start_date_str: str,
    end_date_str: str,
    target_user_id: Optional[str] = None,
) -> Dict[str, Any]:
    """Aggregates attendance data across users and dates into structured report format."""
    user_map = {str(u["_id"]): u for u in users}

    records_in_range = [
        r for r in attendance_records
        if start_date_str <= r.get("date", "") <= end_date_str
    ]

    is_single = bool(target_user_id and target_user_id != "all")
    if is_single:
        records_in_range = [r for r in records_in_range if str(r.get("user_id")) == target_user_id]
        target_u = user_map.get(target_user_id)
        users_in_scope = [target_u] if target_u else []
    else:
        active_uids = set(str(r.get("user_id")) for r in records_in_range)
        users_in_scope = [u for u in users if str(u["_id"]) in active_uids]
        if not users_in_scope:
            users_in_scope = users

    records_by_uid: Dict[str, List[Dict[str, Any]]] = {}
    for r in records_in_range:
        uid = str(r.get("user_id"))
        records_by_uid.setdefault(uid, []).append(r)

    summary_rows = []
    emp_details = {}
    sorted_users = sorted(users_in_scope, key=lambda u: (u.get("name") or "").strip().lower())

    for idx, u in enumerate(sorted_users, 1):
        uid = str(u["_id"])
        raw_name = (u.get("name") or "Unknown").strip()
        raw_role = u.get("role") or "sales"
        masked_role = mask_role(raw_name, raw_role)

        user_recs = records_by_uid.get(uid, [])
        user_recs.sort(key=lambda r: r.get("date", ""))

        present_count = sum(1 for r in user_recs if r.get("status") == "present")
        worked_sec = sum(r.get("worked_seconds", 0) for r in user_recs)
        ot_sec = sum(r.get("overtime_seconds", 0) for r in user_recs)
        late_count = sum(1 for r in user_recs if (r.get("late_seconds", 0) or 0) > 300)

        summary_rows.append({
            "sno": idx,
            "user_id": uid,
            "name": raw_name,
            "role": masked_role,
            "present_days": present_count,
            "worked_seconds": worked_sec,
            "worked_hm": sec_to_hm(worked_sec),
            "overtime_seconds": ot_sec,
            "overtime_hm": sec_to_hm(ot_sec),
            "late_markings": late_count,
        })

        daily_punches = []
        for r in user_recs:
            d_str = r.get("date", "")
            try:
                day_name = datetime.strptime(d_str, "%Y-%m-%d").strftime("%A")
            except Exception:
                day_name = "-"

            late_s = r.get("late_seconds", 0) or 0
            ot_s = r.get("overtime_seconds", 0) or 0
            daily_punches.append({
                "date": d_str,
                "day": day_name,
                "status": (r.get("status") or "present").capitalize(),
                "check_in": to_ist_str(r.get("check_in_at")),
                "check_out": to_ist_str(r.get("check_out_at")),
                "worked_hm": sec_to_hm(r.get("worked_seconds", 0)),
                "overtime_hm": sec_to_hm(ot_s) if ot_s > 0 else "0h 00m",
                "late_mins": round(late_s / 60) if late_s > 300 else 0,
                "location": "WFH" if r.get("check_in_wfh") else "Office",
            })

        emp_details[uid] = {
            "name": raw_name,
            "role": masked_role,
            "punches": daily_punches,
        }

    totals = {
        "present_days": sum(r["present_days"] for r in summary_rows),
        "worked_seconds": sum(r["worked_seconds"] for r in summary_rows),
        "worked_hm": sec_to_hm(sum(r["worked_seconds"] for r in summary_rows)),
        "overtime_seconds": sum(r["overtime_seconds"] for r in summary_rows),
        "overtime_hm": sec_to_hm(sum(r["overtime_seconds"] for r in summary_rows)),
        "late_markings": sum(r["late_markings"] for r in summary_rows),
    }

    cur = datetime.strptime(start_date_str, "%Y-%m-%d")
    end_dt = datetime.strptime(end_date_str, "%Y-%m-%d")
    all_dates = []
    while cur <= end_dt:
        all_dates.append(cur.strftime("%Y-%m-%d"))
        cur += timedelta(days=1)

    try:
        dt_start = datetime.strptime(start_date_str, "%Y-%m-%d")
        dt_end = datetime.strptime(end_date_str, "%Y-%m-%d")
        if dt_start.year == dt_end.year and dt_start.month == dt_end.month:
            period_label = dt_start.strftime("%B %Y")
        else:
            period_label = f"{dt_start.strftime('%d %b %Y')} to {dt_end.strftime('%d %b %Y')}"
    except Exception:
        period_label = f"{start_date_str} to {end_date_str}"

    return {
        "start_date": start_date_str,
        "end_date": end_date_str,
        "period_label": period_label,
        "total_tracked": len(summary_rows),
        "summary_rows": summary_rows,
        "totals": totals,
        "all_dates": all_dates,
        "emp_details": emp_details,
        "is_single_user": is_single,
        "target_user_id": target_user_id,
    }


def generate_attendance_excel_bytes(report_data: Dict[str, Any]) -> bytes:
    """Generates the multi-sheet Excel workbook as bytes."""
    wb = openpyxl.Workbook()

    navy_fill = PatternFill("solid", fgColor="1E293B")
    header_fill = PatternFill("solid", fgColor="F1F5F9")
    alt_fill = PatternFill("solid", fgColor="F8FAFC")
    present_fill = PatternFill("solid", fgColor="DCFCE7")
    absent_fill = PatternFill("solid", fgColor="FEE2E2")

    font_title = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    font_head = Font(name="Arial", size=10, bold=True, color="334155")
    font_bold = Font(name="Arial", size=9, bold=True, color="1E293B")
    font_reg = Font(name="Arial", size=9, color="0F172A")
    font_present = Font(name="Arial", size=9, bold=True, color="15803D")
    font_absent = Font(name="Arial", size=9, color="B91C1C")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    period_label = report_data.get("period_label", "Report")

    # SHEET 1: Monthly Summary
    ws_sum = wb.active
    ws_sum.title = "Monthly Summary"

    ws_sum.merge_cells("A1:G1")
    t_cell = ws_sum["A1"]
    t_cell.value = f"UNIQUE PRIME REALITY — ATTENDANCE SUMMARY ({period_label.upper()})"
    t_cell.font = font_title
    t_cell.fill = navy_fill
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 32

    headers = ["S.No", "Employee Name", "Role", "Days Present", "Total Hours Worked", "Total Overtime", "Late Markings"]
    ws_sum.append([])
    ws_sum.append(headers)
    ws_sum.row_dimensions[3].height = 24

    for c_idx in range(1, len(headers) + 1):
        c = ws_sum.cell(row=3, column=c_idx)
        c.fill = header_fill
        c.font = font_head
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", vertical="center")

    r_idx = 4
    for row in report_data.get("summary_rows", []):
        ws_sum.append([
            row["sno"],
            row["name"],
            row["role"],
            row["present_days"],
            row["worked_hm"],
            row["overtime_hm"],
            row["late_markings"],
        ])
        ws_sum.row_dimensions[r_idx].height = 20
        for c_idx in range(1, 8):
            c = ws_sum.cell(row=r_idx, column=c_idx)
            c.font = font_reg
            c.border = thin_border
            if c_idx in [1, 3, 4, 5, 6, 7]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        r_idx += 1

    totals = report_data.get("totals", {})
    ws_sum.append([
        "TOTAL",
        f"{report_data.get('total_tracked', 0)} Employees",
        "—",
        totals.get("present_days", 0),
        totals.get("worked_hm", "0h 00m"),
        totals.get("overtime_hm", "0h 00m"),
        totals.get("late_markings", 0),
    ])
    ws_sum.row_dimensions[r_idx].height = 22
    for c_idx in range(1, 8):
        c = ws_sum.cell(row=r_idx, column=c_idx)
        c.font = font_bold
        c.fill = header_fill
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", vertical="center")

    for col in ws_sum.columns:
        max_len = max(len(str(cell.value or "")) for cell in col if cell.row > 1)
        ws_sum.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)
    ws_sum.column_dimensions["B"].width = 24

    # SHEET 2: Attendance Matrix (Multi-user only)
    if not report_data.get("is_single_user"):
        ws_matrix = wb.create_sheet(title="Attendance Matrix")
        all_dates = report_data.get("all_dates", [])

        date_headers = ["Employee Name"]
        for d in all_dates:
            dt = datetime.strptime(d, "%Y-%m-%d")
            date_headers.append(f"{dt.day} {dt.strftime('%b')}")
        date_headers.append("Total Present")

        ws_matrix.append(date_headers)
        ws_matrix.row_dimensions[1].height = 25
        for c_idx in range(1, len(date_headers) + 1):
            c = ws_matrix.cell(row=1, column=c_idx)
            c.fill = navy_fill
            c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border

        m_idx = 2
        for row in report_data.get("summary_rows", []):
            uid = row["user_id"]
            emp_info = report_data.get("emp_details", {}).get(uid, {})
            punch_dates = {p["date"]: p for p in emp_info.get("punches", [])}

            m_row = [row["name"]]
            pres_count = 0
            for d in all_dates:
                p = punch_dates.get(d)
                dt = datetime.strptime(d, "%Y-%m-%d")
                if p:
                    m_row.append("P")
                    pres_count += 1
                elif dt.weekday() == 6:
                    m_row.append("OFF")
                else:
                    m_row.append("A")
            m_row.append(pres_count)
            ws_matrix.append(m_row)
            ws_matrix.row_dimensions[m_idx].height = 18

            ws_matrix.cell(row=m_idx, column=1).font = font_bold
            ws_matrix.cell(row=m_idx, column=1).border = thin_border

            for c_idx in range(2, len(all_dates) + 2):
                cell = ws_matrix.cell(row=m_idx, column=c_idx)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                val = cell.value
                if val == "P":
                    cell.fill = present_fill
                    cell.font = font_present
                elif val == "OFF":
                    cell.fill = alt_fill
                    cell.font = Font(name="Arial", size=8, color="94A3B8")
                elif val == "A":
                    cell.fill = absent_fill
                    cell.font = font_absent

            tot_c = ws_matrix.cell(row=m_idx, column=len(all_dates) + 2)
            tot_c.font = font_bold
            tot_c.fill = header_fill
            tot_c.border = thin_border
            tot_c.alignment = Alignment(horizontal="center", vertical="center")
            m_idx += 1

        ws_matrix.column_dimensions["A"].width = 22
        for c_idx in range(2, len(all_dates) + 2):
            ws_matrix.column_dimensions[get_column_letter(c_idx)].width = 7
        ws_matrix.column_dimensions[get_column_letter(len(all_dates) + 2)].width = 14

    # INDIVIDUAL EMPLOYEE TABS
    for row in report_data.get("summary_rows", []):
        uid = row["user_id"]
        emp_info = report_data.get("emp_details", {}).get(uid, {})
        punches = emp_info.get("punches", [])

        tab_name = sanitize_sheet_title(row["name"])
        ws_emp = wb.create_sheet(title=tab_name)

        ws_emp.merge_cells("A1:H1")
        h_cell = ws_emp["A1"]
        h_cell.value = f"{row['name'].upper()} — ATTENDANCE STATEMENT ({period_label.upper()})"
        h_cell.font = font_title
        h_cell.fill = navy_fill
        h_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_emp.row_dimensions[1].height = 28

        ws_emp.merge_cells("A2:B2")
        ws_emp["A2"].value = f"Designation: {row['role']}"
        ws_emp["A2"].font = font_bold

        ws_emp.merge_cells("C2:D2")
        ws_emp["C2"].value = f"Days Present: {row['present_days']}"
        ws_emp["C2"].font = font_bold

        ws_emp.merge_cells("E2:F2")
        ws_emp["E2"].value = f"Hours Worked: {row['worked_hm']}"
        ws_emp["E2"].font = font_bold

        ws_emp.merge_cells("G2:H2")
        ws_emp["G2"].value = f"Overtime: {row['overtime_hm']}"
        ws_emp["G2"].font = font_bold
        ws_emp.row_dimensions[2].height = 20

        emp_headers = ["Date", "Day", "Status", "Check-in (IST)", "Check-out (IST)", "Hours Worked", "Late (Mins)", "Location"]
        ws_emp.append([])
        ws_emp.append(emp_headers)
        ws_emp.row_dimensions[4].height = 22

        for c_idx in range(1, len(emp_headers) + 1):
            c = ws_emp.cell(row=4, column=c_idx)
            c.fill = header_fill
            c.font = font_head
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")

        p_row_idx = 5
        for p in punches:
            ws_emp.append([
                p["date"],
                p["day"],
                p["status"],
                p["check_in"],
                p["check_out"],
                p["worked_hm"],
                p["late_mins"] if p["late_mins"] > 0 else "—",
                p["location"],
            ])
            ws_emp.row_dimensions[p_row_idx].height = 18
            for c_idx in range(1, 9):
                c = ws_emp.cell(row=p_row_idx, column=c_idx)
                c.font = font_reg
                c.border = thin_border
                c.alignment = Alignment(horizontal="center", vertical="center")
            p_row_idx += 1

        for col in ws_emp.columns:
            if col[0].row >= 4:
                max_len = max(len(str(cell.value or "")) for cell in col if cell.row >= 4)
                ws_emp.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_attendance_pdf_bytes(report_data: Dict[str, Any]) -> bytes:
    """Generates the styled corporate PDF report as bytes."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    BRAND = colors.HexColor("#1a3fbf")
    SLATE_900 = colors.HexColor("#0f172a")
    SLATE_800 = colors.HexColor("#1e293b")
    SLATE_400 = colors.HexColor("#94a3b8")
    SLATE_100 = colors.HexColor("#f1f5f9")
    BORDER_COLOR = colors.HexColor("#cbd5e1")

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, textColor=SLATE_900, leading=22)
    sub_style = ParagraphStyle("sub", fontName="Helvetica", fontSize=9, textColor=SLATE_400, leading=12)
    brand_style = ParagraphStyle("brand", fontName="Helvetica-Bold", fontSize=11, textColor=BRAND, leading=14)
    cell_head = ParagraphStyle("cell_head", fontName="Helvetica-Bold", fontSize=8.5, textColor=colors.white, alignment=TA_CENTER)
    cell_body = ParagraphStyle("cell_body", fontName="Helvetica", fontSize=8, textColor=SLATE_800, alignment=TA_CENTER)
    cell_name = ParagraphStyle("cell_name", fontName="Helvetica-Bold", fontSize=8.5, textColor=SLATE_900, alignment=TA_LEFT)
    cell_total = ParagraphStyle("cell_total", fontName="Helvetica-Bold", fontSize=8.5, textColor=SLATE_900, alignment=TA_CENTER)

    elements = []
    period_label = report_data.get("period_label", "Report")

    # Header Banner
    header_table = Table([
        [
            Paragraph("UNIQUE PRIME REALITY", brand_style),
            Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y')}", sub_style)
        ],
        [
            Paragraph(f"Monthly Attendance Report — {period_label}", title_style),
            Paragraph(f"Tracked: {report_data.get('total_tracked', 0)} members", sub_style)
        ]
    ], colWidths=[130 * mm, 60 * mm])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(header_table)
    elements.append(Spacer(1, 4 * mm))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=BRAND, spaceAfter=8))

    # KPI Summary Boxes
    totals = report_data.get("totals", {})
    kpi_data = [
        [
            Paragraph("<b>TOTAL MEMBERS</b><br/>" + str(report_data.get("total_tracked", 0)), cell_total),
            Paragraph("<b>TOTAL PRESENT</b><br/>" + str(totals.get("present_days", 0)) + " Days", cell_total),
            Paragraph("<b>TOTAL HOURS</b><br/>" + totals.get("worked_hm", "0h 00m"), cell_total),
            Paragraph("<b>TOTAL OVERTIME</b><br/>" + totals.get("overtime_hm", "0h 00m"), cell_total),
            Paragraph("<b>LATE MARKINGS</b><br/>" + str(totals.get("late_markings", 0)), cell_total),
        ]
    ]
    kpi_table = Table(kpi_data, colWidths=[38 * mm] * 5)
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SLATE_100),
        ("BOX", (0, 0), (-1, -1), 1, BORDER_COLOR),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 6 * mm))

    # Main Summary Table
    table_headers = ["S.No", "Employee Name", "Role", "Days Present", "Hours Worked", "Overtime", "Late Marks"]
    col_widths = [12 * mm, 45 * mm, 32 * mm, 26 * mm, 28 * mm, 25 * mm, 22 * mm]
    t_rows = [[Paragraph(h, cell_head) for h in table_headers]]

    for row in report_data.get("summary_rows", []):
        t_rows.append([
            Paragraph(str(row["sno"]), cell_body),
            Paragraph(row["name"], cell_name),
            Paragraph(row["role"], cell_body),
            Paragraph(str(row["present_days"]), cell_body),
            Paragraph(row["worked_hm"], cell_body),
            Paragraph(row["overtime_hm"], cell_body),
            Paragraph(str(row["late_markings"]), cell_body),
        ])

    t_rows.append([
        Paragraph("—", cell_total),
        Paragraph("<b>TOTAL / OVERALL</b>", cell_name),
        Paragraph("—", cell_total),
        Paragraph(f"<b>{totals.get('present_days', 0)}</b>", cell_total),
        Paragraph(f"<b>{totals.get('worked_hm', '0h 00m')}</b>", cell_total),
        Paragraph(f"<b>{totals.get('overtime_hm', '0h 00m')}</b>", cell_total),
        Paragraph(f"<b>{totals.get('late_markings', 0)}</b>", cell_total),
    ])

    summary_table = Table(t_rows, colWidths=col_widths, repeatRows=1)
    ts = [
        ("BACKGROUND", (0, 0), (-1, 0), SLATE_800),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
        ("BACKGROUND", (0, -1), (-1, -1), SLATE_100),
    ]
    for i in range(1, len(t_rows) - 1):
        if i % 2 == 0:
            ts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f8fafc")))

    summary_table.setStyle(TableStyle(ts))
    elements.append(summary_table)

    # If single user: Day-by-Day Punch Statement
    if report_data.get("is_single_user") and report_data.get("summary_rows"):
        single_uid = report_data["summary_rows"][0]["user_id"]
        emp_punches = report_data.get("emp_details", {}).get(single_uid, {}).get("punches", [])
        if emp_punches:
            elements.append(Spacer(1, 6 * mm))
            p_head = Paragraph(f"<b>Daily Punch Log & In/Out Timings — {report_data['summary_rows'][0]['name']}</b>", title_style)
            elements.append(p_head)
            elements.append(Spacer(1, 3 * mm))

            log_headers = ["Date", "Day", "Status", "Check-in (IST)", "Check-out (IST)", "Worked", "Overtime", "Late"]
            log_widths = [22 * mm, 24 * mm, 20 * mm, 26 * mm, 26 * mm, 25 * mm, 24 * mm, 23 * mm]
            log_rows = [[Paragraph(h, cell_head) for h in log_headers]]

            for p in emp_punches:
                log_rows.append([
                    Paragraph(p["date"], cell_body),
                    Paragraph(p["day"], cell_body),
                    Paragraph(p["status"], cell_body),
                    Paragraph(p["check_in"], cell_body),
                    Paragraph(p["check_out"], cell_body),
                    Paragraph(p["worked_hm"], cell_body),
                    Paragraph(p["overtime_hm"], cell_body),
                    Paragraph(str(p["late_mins"]) + "m" if p["late_mins"] > 0 else "—", cell_body),
                ])

            log_table = Table(log_rows, colWidths=log_widths, repeatRows=1)
            lts = [
                ("BACKGROUND", (0, 0), (-1, 0), SLATE_800),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, BORDER_COLOR),
                ("TOPPADDING", (0, 0), (-1, -1), 3.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ]
            for i in range(1, len(log_rows)):
                if i % 2 == 0:
                    lts.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f8fafc")))
            log_table.setStyle(TableStyle(lts))
            elements.append(log_table)

    elements.append(Spacer(1, 8 * mm))
    footer_p = Paragraph(
        "<i>Confidential attendance record generated by Unique Prime Reality CRM. All timestamps converted to Indian Standard Time (IST).</i>",
        ParagraphStyle("ft", fontName="Helvetica", fontSize=7.5, textColor=SLATE_400, alignment=TA_CENTER),
    )
    elements.append(footer_p)

    doc.build(elements)
    return buf.getvalue()
